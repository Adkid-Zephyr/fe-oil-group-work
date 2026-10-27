"""Reproducible, non-imputing cleaning of the complete course oil workbook.

Only formatting and deterministic derivations are changed. Source anomalies are
retained or quarantined with exact Excel locations. Run from any directory.
"""
from pathlib import Path
from datetime import datetime, date
import argparse
import hashlib
import json
import re
import numpy as np
import pandas as pd
import openpyxl

ROOT = Path(__file__).resolve().parent
REPOSITORY = ROOT.parent
RAW = REPOSITORY / 'data' / 'raw' / 'Oil Data.xlsx'
EXPECTED_SHA = '4355813a4761a6f242ae6ff1eff03dce6f89127320161b6fb2481e9aa908bae5'
# Names distinguish quantities whose source label is misleading or incomplete.
SCHEMAS = {
 'Industry_D': ('industry_daily.csv', ['cl1','cl2','brent','crack_spread_321','baltic_dry_index']),
 'Industry_W': ('industry_weekly.csv', ['crude_production_source_units','rig_count','crude_stock_change_source_units','commercial_long','commercial_short','total_open_interest_source']),
 'Industry_M': ('industry_monthly.csv', ['cl1','cl2','brent','crack_spread_321','baltic_dry_index','crude_production_source_units','inventories_source_units','rig_count','commercial_long','commercial_short','commercial_long_plus_short_source']),
 'Macro_M': ('macro_monthly.csv', ['cpi_yoy_pct','treasury_3m_yield_pct','treasury_10y_yield_pct','industrial_production_index','sp500_index','dxy_return_decimal','vix']),
 'StockPrices_M': ('stock_prices_monthly.csv', ['gasoline_all_grades','dow_jones_us_oil_gas_index','exxon_mobil','conoco_phillips','chevron','bp','shell','total_energies']),
}

def digest(p):
 return hashlib.sha256(Path(p).read_bytes()).hexdigest()

def normalize_records(rows, sheet, names):
 """Validate before sorting; never guess dates or silently coerce bad numbers."""
 header = rows[0]
 n = len(names)+1
 assert len(header)>=n and all(v is not None for v in header[:n]), f'{sheet}: incomplete header'
 assert all(v is None for v in header[n:]), f'{sheet}: unexpected extra headings'
 records, quarantine, empty_rows = [], [], []
 for excel_row, row in enumerate(rows[1:],start=2):
  if all(v is None for v in row):
   empty_rows.append(excel_row);continue
  assert all(v is None for v in row[n:]), f'{sheet}!{excel_row}: values in unnamed columns'
  values={}
  for name,v in zip(names,row[1:n]):
   if v is not None:
    assert isinstance(v,(int,float)) and not isinstance(v,bool) and np.isfinite(v), f'{sheet}!{excel_row}: nonnumeric {name}={v!r}'
   values[name]=v
  if row[0] is None:
   quarantine.append({'source_sheet':sheet,'source_excel_row':excel_row,'reason':'missing_date','raw_values_json':json.dumps(dict(zip([str(h).strip() for h in header[:n]],row[:n])),ensure_ascii=False)})
   continue
  assert isinstance(row[0],(datetime,date)), f'{sheet}!A{excel_row}: unexpected date {row[0]!r}'
  stamp=pd.Timestamp(row[0])
  assert stamp==stamp.normalize(), f'{sheet}: intraday timestamp requires review'
  records.append(dict(date=stamp,source_excel_row=excel_row,**values))
 df=pd.DataFrame(records)
 assert not df.date.duplicated().any(), f'{sheet}: duplicate dates need explicit resolution'
 df=df.sort_values('date').reset_index(drop=True)
 if sheet.endswith('_M'):
  df.insert(1,'month',df.date.dt.strftime('%Y-%m'))
  assert not df.month.duplicated().any(), f'{sheet}: duplicate months'
 return df,quarantine,dict(raw_rows=len(rows)-1,dated_rows=len(df),blank_rows=len(empty_rows),quarantined_rows=len(quarantine),removed_blank_columns=len(header)-n,blank_excel_rows=empty_rows)

def make_targets(monthly):
 x=monthly.set_index('month').copy()
 grid=pd.period_range(x.index.min(),x.index.max(),freq='M').astype(str)
 x=x.reindex(grid)
 price=x.cl1;previous=price.shift(1)
 valid=price.notna() & previous.notna() & (price>0) & (previous>0)
 # Nonpositive endpoints are not discarded from native data; derived target is
 # undefined under this conventional positive-price percentage-change protocol.
 out=pd.DataFrame({'month':grid,'origin':(pd.PeriodIndex(grid,freq='M')-1).astype(str),
  'actual_price':price.to_numpy(),'benchmark_price':previous.to_numpy(),
  'actual_return_pct':(100*(price/previous-1)).where(valid).to_numpy(),
  'benchmark_return_pct':np.where(valid,0.,np.nan),'target_valid':valid.to_numpy(),
  'is_shared_evaluation':(grid>='2015-01')&(grid<='2024-06')})
 return out

def write_csv(frame,path):
 frame.to_csv(path,index=False,date_format='%Y-%m-%d',float_format='%.17g',na_rep='')

def build(raw=RAW,out_root=REPOSITORY):
 raw=Path(raw);out_root=Path(out_root)
 before=digest(raw)
 assert before==EXPECTED_SHA, 'Source hash changed: inspect and version the new source before rerunning'
 clean=out_root/'data'/'clean';qa=out_root/'guidelines'/'Data quality'
 clean.mkdir(parents=True,exist_ok=True);qa.mkdir(parents=True,exist_ok=True)
 cached=openpyxl.load_workbook(raw,read_only=True,data_only=True)
 formula_book=openpyxl.load_workbook(raw,read_only=True,data_only=False)
 assert set(cached.sheetnames)=={'Data Definitions',*SCHEMAS}, 'Unexpected workbook sheets'
 definitions=[]
 for i,row in enumerate(cached['Data Definitions'].values,start=1):
  if any(v is not None for v in row):definitions.append(dict(source_excel_row=i,label=row[0],description=row[1]))
 write_csv(pd.DataFrame(definitions),qa/'source_definitions.csv')
 frames={};summary={};quarantined=[];dictionary=[];formulas=[];missing=[];gaps=[];flags=[];runs=[];missing_ranges=[]
 for sheet,(filename,names) in SCHEMAS.items():
  rows=list(cached[sheet].values)
  df,quarantine,stats=normalize_records(rows,sheet,names)
  quarantined.extend(quarantine)
  stats.update(start=str(df.date.min().date()),end=str(df.date.max().date()),original_order='descending' if pd.Series([r[0] for r in rows[1:] if r[0] is not None]).is_monotonic_decreasing else 'other',duplicate_dates=0)
  summary[sheet]=stats
  frames[sheet]=df
  for col,name in enumerate(names,start=2):
   original=str(rows[0][col-1]).strip()
   status='retained_as_supplied';note='Observation/reference date, not verified publication or execution time.'
   unit='Source numeric units; exact scale/adjustment not independently verified'
   if name in ['cl1','cl2','brent']:unit='USD per barrel; continuous futures/benchmark construction not documented'
   if sheet=='Industry_D' and name in ['brent','crack_spread_321','baltic_dry_index']:
    status='alignment_review_required';note='Large same-date monthly discrepancies; daily sheet also has undated trailing Brent values. Do not assume joint-calendar alignment is validated.'
   if sheet=='Industry_M' and name in ['brent','crack_spread_321','baltic_dry_index']:
    status='cross_frequency_review_required';note='Same-date daily/monthly values frequently differ. Aggregation, contract and timestamp conventions are unresolved; do not substitute one series for the other.'
   if sheet=='Industry_M' and name=='commercial_long_plus_short_source':
    status='misleading_source_label';note='Source header is Total Open Interest but all 355 cells are formulas J+K (commercial long + short). Preserved under a literal descriptive name. Not independently measured total open interest.'
   if sheet=='Industry_W' and name=='commercial_long':
    status='contains_source_interpolation';note='Source E1195, dated 2001-09-07, is the average of adjacent source rows. Uses both an earlier and a later date: look-ahead risk for real-time forecasting. Value retained; use commercial_long_unimputed to mask it.'
   if sheet=='StockPrices_M':
    note+=' Stock split/dividend adjustment, listing/currency conventions not documented.'
   if name=='dxy_return_decimal':unit='Decimal return (working interpretation; raw scale retained)'
   if name.endswith('_pct'):unit='Percentage points as supplied; no additional multiplication'
   dictionary.append(dict(source_sheet=sheet,source_column=original,clean_file=filename,clean_column=name,excel_column=openpyxl.utils.get_column_letter(col),unit=unit,status=status,note=note))
   series=df[name];present=df.loc[series.notna(),'date']
   first=present.min() if len(present) else pd.NaT;last=present.max() if len(present) else pd.NaT
   mask=series.isna()
   missing.append(dict(sheet=sheet,column=name,n_rows=len(df),n_missing=int(mask.sum()),missing_pct=float(mask.mean()*100),first_valid=first,last_valid=last,leading_missing=int((mask&(df.date<first)).sum()),internal_missing=int((mask&(df.date>=first)&(df.date<=last)).sum()),trailing_missing=int((mask&(df.date>last)).sum())))
   for _,ix in df.groupby(mask.ne(mask.shift()).cumsum()).groups.items():
    if bool(mask.loc[ix[0]]):
     missing_ranges.append(dict(sheet=sheet,column=name,start=df.loc[ix[0],'date'],end=df.loc[ix[-1],'date'],n_observations=len(ix),action='left missing'))
   # Descriptive full-sample outlier flags, never filtering or model features.
   delta=series.diff();med=delta.median();mad=(delta-med).abs().median()
   if pd.notna(mad) and mad>0:
    score=.6744897501960817*(delta-med)/mad
    for j in df.index[score.abs()>8]:
     flags.append(dict(sheet=sheet,date=df.loc[j,'date'],source_excel_row=int(df.loc[j,'source_excel_row']),column=name,value=float(series[j]),kind='large_first_difference_mad',detail=f'abs(modified_z)>8; z={score[j]:.6g}; descriptive only, retained'))
   for j in df.index[series<0]:
    flags.append(dict(sheet=sheet,date=df.loc[j,'date'],source_excel_row=int(df.loc[j,'source_excel_row']),column=name,value=float(series[j]),kind='negative_value_review',detail='Retained. Negative prices, changes, spreads or rates may be economically meaningful.'))
   # Consecutive repeats can be stale quotes or legitimate observations.
   group=(series.ne(series.shift()) | series.isna()).cumsum()
   threshold=5 if sheet.endswith('_D') else 4 if sheet.endswith('_W') else 3
   for _,ix in df.groupby(group).groups.items():
    if len(ix)>=threshold and pd.notna(series.loc[ix[0]]):
     runs.append(dict(sheet=sheet,column=name,start=df.loc[ix[0],'date'],end=df.loc[ix[-1],'date'],observations=len(ix),value=float(series.loc[ix[0]]),action='retained; review only'))
  if sheet.endswith('_M'):
   expected=set(pd.period_range(df.month.min(),df.month.max(),freq='M').astype(str));absent=sorted(expected-set(df.month))
   assert not absent, f'{sheet}: missing months {absent}'
   stats['missing_months']=0
  elif sheet.endswith('_W'):
   days=df.date.diff().dt.days;stats['non_7_day_steps']=int((days.dropna()!=7).sum())
  else:
   # Business days are not an exchange holiday calendar. List for review only.
   absent=pd.bdate_range(df.date.min(),df.date.max()).difference(df.date)
   for d in absent:gaps.append(dict(sheet=sheet,date=d,classification='weekday_without_row',action='not imputed; may be exchange holiday or missing source observation'))
   stats['unverified_weekday_gaps']=len(absent);stats['weekend_rows']=int((df.date.dt.day_of_week>4).sum())
  # Enumerate all original formulas without evaluating arbitrary expressions.
  for row in formula_book[sheet]:
   for cell in row:
    assert cell.data_type!='e', f'Excel error {sheet}!{cell.coordinate}'
    if cell.data_type!='f':continue
    cached_value=rows[cell.row-1][cell.column-1]
    assert isinstance(cached_value,(int,float)), f'Missing formula cache {sheet}!{cell.coordinate}'
    expression=cell.value;compact=expression.replace(' ','')
    match=re.fullmatch(r'=J(\d+)\+K\1',compact)
    interp=re.fullmatch(r'=\(E(\d+)\+E(\d+)\)/2',compact)
    if sheet=='Industry_M' and match:
     rid=int(match[1]);expected_value=rows[rid-1][9]+rows[rid-1][10];kind='commercial_long_plus_short'
    elif sheet=='Industry_W' and interp:
     expected_value=(rows[int(interp[1])-1][4]+rows[int(interp[2])-1][4])/2;kind='source_linear_interpolation'
    else:raise AssertionError(f'Unreviewed formula: {sheet}!{cell.coordinate} {expression}')
    assert abs(cached_value-expected_value)<1e-9,'Formula cache mismatch'
    formulas.append(dict(sheet=sheet,cell=cell.coordinate,source_excel_row=cell.row,date=rows[cell.row-1][0],formula=expression,cached_value=cached_value,recomputed_value=expected_value,kind=kind,action='retained; cache verified'))
  if sheet=='Macro_M':
   df['dxy_return_pct']=df.dxy_return_decimal*100
   dictionary.append(dict(source_sheet=sheet,source_column='DXY returns',clean_file=filename,clean_column='dxy_return_pct',excel_column='G',unit='Percentage-point return, derived as 100 * dxy_return_decimal',status='derived_scale_conversion',note='Decimal interpretation is documented, not inferred from Excel percent formatting (source cells use General).'))
  if sheet=='Industry_W':
   df['commercial_long_is_source_interpolated']=df.source_excel_row.eq(1195)
   df['commercial_long_unimputed']=df.commercial_long.mask(df.commercial_long_is_source_interpolated)
   dictionary.append(dict(source_sheet=sheet,source_column='Weekly Commercial Long',clean_file=filename,clean_column='commercial_long_unimputed',excel_column='E',unit='Contracts as supplied',status='derived_source_interpolation_mask',note='Same as commercial_long except source-interpolated E1195 is missing; avoids relying on its future endpoint. Original value retained in commercial_long.'))
  write_csv(df,clean/filename)
 # Monthly observation panel: no lags, resampling, imputation or release-date claims.
 panels=[]
 for sheet,prefix in [('Industry_M','industry'),('Macro_M','macro'),('StockPrices_M','stocks')]:
  part=frames[sheet].set_index('month').add_prefix(prefix+'__');panels.append(part)
 joined=pd.concat(panels,axis=1,join='outer').sort_index().rename_axis('month').reset_index()
 assert len(joined)==355 and not joined.month.duplicated().any()
 write_csv(joined,clean/'monthly_observations.csv')
 targets=make_targets(frames['Industry_M']);write_csv(targets,clean/'monthly_targets.csv')
 # Compare shared targets against the existing repository benchmark, unchanged.
 prior=out_root/'data'/'results'/'no-change-benchmark'/'forecasts.csv'
 benchmark_check='not_available'
 if prior.exists():
  a=targets.loc[targets.is_shared_evaluation].copy();b=pd.read_csv(prior)
  assert list(a.month)==list(b.month) and len(a)==114
  for col in ['actual_price','benchmark_price','actual_return_pct','benchmark_return_pct']:
   assert np.allclose(a[col],b[col],rtol=1e-12,atol=1e-12),f'Changed shared benchmark {col}'
  benchmark_check='matches all 114 rows'
 # Same-date cross-frequency comparison. Differences do not prove which sheet is wrong.
 daily=frames['Industry_D'].set_index('date');monthly=frames['Industry_M'];comparisons=[]
 for row in monthly.itertuples():
  for col in ['cl1','cl2','brent','crack_spread_321','baltic_dry_index']:
   mv=getattr(row,col);dv=daily.loc[row.date,col] if row.date in daily.index else np.nan
   status='daily_date_or_value_absent' if pd.isna(dv) else 'match' if abs(mv-dv)<1e-8 else 'different'
   comparisons.append(dict(date=row.date,column=col,monthly_value=mv,daily_value=dv,difference=mv-dv if pd.notna(dv) else np.nan,status=status,monthly_source_excel_row=row.source_excel_row,daily_source_excel_row=daily.loc[row.date,'source_excel_row'] if row.date in daily.index else np.nan))
 comp=pd.DataFrame(comparisons);write_csv(comp,qa/'cross_frequency_comparison.csv')
 cross=comp.groupby(['column','status']).size().unstack(fill_value=0).reset_index();write_csv(cross,qa/'cross_frequency_summary.csv')
 write_csv(pd.DataFrame(quarantined,columns=['source_sheet','source_excel_row','reason','raw_values_json']),qa/'quarantined_rows.csv')
 write_csv(pd.DataFrame(missing),qa/'missingness.csv')
 write_csv(pd.DataFrame(missing_ranges),qa/'missing_ranges.csv')
 write_csv(pd.DataFrame(flags,columns=['sheet','date','source_excel_row','column','value','kind','detail']),qa/'review_flags.csv')
 write_csv(pd.DataFrame(runs,columns=['sheet','column','start','end','observations','value','action']),qa/'constant_runs.csv')
 write_csv(pd.DataFrame(gaps,columns=['sheet','date','classification','action']),qa/'calendar_review.csv')
 write_csv(pd.DataFrame(formulas),qa/'formula_audit.csv')
 write_csv(pd.DataFrame(dictionary),qa/'data_dictionary.csv')
 assert digest(raw)==before,'Raw workbook was modified'
 summary_json=dict(dataset_version='1.0.0',source_sha256=before,source_name='data/raw/Oil Data.xlsx',source_immutable=True,script_sha256=digest(__file__),sheets=summary,quarantined_rows=len(quarantined),formulas_reviewed=len(formulas),review_flags=len(flags),constant_runs=len(runs),monthly_rows=len(joined),shared_target_rows=int(targets.is_shared_evaluation.sum()),shared_benchmark_check=benchmark_check,imputed_cells_by_this_pipeline=0,deleted_numeric_values=0,unresolved=['Non-CL daily series alignment and cross-frequency definition mismatches','Original contract roll/adjustment methodology','Publication timestamps and historical revisions','Production/inventory scales and some instrument metadata','Whether source interpolation is acceptable for model estimation'])
 (qa/'summary.json').write_text(json.dumps(summary_json,indent=2)+'\n')
 files={str(p.relative_to(out_root)):digest(p) for folder in [clean,qa] for p in sorted(folder.glob('*')) if p.suffix in ['.csv','.json'] and p.name!='manifest.json'}
 (qa/'manifest.json').write_text(json.dumps({'source_sha256':before,'files_sha256':files},indent=2)+'\n')
 print(json.dumps({k:v for k,v in summary_json.items() if k!='sheets'},indent=2))

if __name__=='__main__':
 parser=argparse.ArgumentParser();parser.add_argument('--output-root',type=Path,default=REPOSITORY);args=parser.parse_args();build(out_root=args.output_root)
